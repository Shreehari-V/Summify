# backend/processing_service.py

import os
import re
import json
import uuid
import asyncio
import subprocess
import tempfile
from pathlib import Path
from datetime import datetime
from typing import Tuple, Dict, Any, Optional, List
from bson import ObjectId
import httpx
import pymupdf
import docx
import imageio_ffmpeg
from sklearn.feature_extraction.text import TfidfVectorizer, ENGLISH_STOP_WORDS

from .config import settings

# Supported extensions
DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".md"}
AUDIO_EXTENSIONS = {".mp3", ".wav", ".m4a"}
VIDEO_EXTENSIONS = {".mp4", ".webm"}

HF_ROUTER_URL = "https://router.huggingface.co/hf-inference/models"
HF_LEGACY_URL = "https://api-inference.huggingface.co/models"
HF_CHAT_COMPLETIONS_URL = "https://router.huggingface.co/v1/chat/completions"

ADDITIONAL_STOPWORDS = {
    "slide", "slides", "lecture", "chapter", "today", "page", "pages", 
    "thank", "thanks", "dr", "professor", "course", "class", "topic", 
    "hello", "welcome", "week", "next", "previous", "exam", "homework",
    "example", "summary", "overview", "section", "part", "discussed",
    "thing", "things", "going", "know", "really", "want", "like", "just"
}
CUSTOM_STOP_WORDS = list(ENGLISH_STOP_WORDS.union(ADDITIONAL_STOPWORDS))


def extract_text_from_pdf(file_path: Path) -> Tuple[str, Dict[str, Any]]:
    """Extract clean text and page count from a PDF using PyMuPDF."""
    doc = pymupdf.open(str(file_path))
    pages_text = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text().strip()
        if text:
            pages_text.append(text)
    
    full_text = "\n\n".join(pages_text)
    metadata = {
        "extractor": "PyMuPDF",
        "total_pages": len(doc),
        "pages_with_text": len(pages_text),
    }
    doc.close()
    return full_text, metadata


def extract_text_from_docx(file_path: Path) -> Tuple[str, Dict[str, Any]]:
    """Extract text from a DOCX file including paragraphs and tables."""
    doc = docx.Document(str(file_path))
    parts = []
    
    # Paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
            
    # Tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)
                
    full_text = "\n\n".join(parts)
    metadata = {
        "extractor": "python-docx",
        "paragraphs_count": len(doc.paragraphs),
        "tables_count": len(doc.tables),
    }
    return full_text, metadata


def extract_text_from_txt(file_path: Path) -> Tuple[str, Dict[str, Any]]:
    """Directly read a plain text or markdown file with multi-encoding fallback."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    full_text = ""
    used_encoding = "utf-8"
    
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                full_text = f.read()
            used_encoding = enc
            break
        except (UnicodeDecodeError, LookupError):
            continue
            
    metadata = {
        "extractor": "direct_read",
        "encoding": used_encoding,
        "line_count": len(full_text.splitlines()),
    }
    return full_text, metadata


def extract_audio_from_video(video_path: Path) -> Path:
    """Extract audio from video file to a temporary MP3 file using bundled imageio-ffmpeg."""
    ffmpeg_exe = imageio_ffmpeg.get_ffmpeg_exe()
    temp_dir = Path(tempfile.gettempdir())
    temp_audio_path = temp_dir / f"summify_audio_{video_path.stem}_{int(datetime.utcnow().timestamp())}.mp3"
    
    cmd = [
        ffmpeg_exe,
        "-y",
        "-i", str(video_path),
        "-vn",
        "-acodec", "libmp3lame",
        "-ar", "16000",
        "-ac", "1",
        "-b:a", "64k",
        str(temp_audio_path),
    ]
    
    proc = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if proc.returncode != 0:
        raise RuntimeError(f"Audio extraction from video failed: {proc.stderr[:300]}")
        
    if not temp_audio_path.exists() or temp_audio_path.stat().st_size == 0:
        raise RuntimeError("Extracted audio file is missing or empty")
        
    return temp_audio_path


async def transcribe_audio_hf(audio_path: Path) -> Tuple[str, Dict[str, Any]]:
    """Send audio file to Hugging Face hosted inference using openai/whisper-large-v3-turbo."""
    if not settings.hf_token or not settings.hf_token.strip():
        raise ValueError(
            "Hugging Face API token (HF_TOKEN) is not configured in .env. "
            "Please add HF_TOKEN to your .env file to enable Whisper audio/video transcription."
        )

    model_name = settings.hf_whisper_model or "openai/whisper-large-v3-turbo"
    endpoints = [
        f"{HF_ROUTER_URL}/{model_name}",
        f"{HF_LEGACY_URL}/{model_name}",
    ]
    
    mime_map = {
        ".mp3": "audio/mpeg",
        ".wav": "audio/wav",
        ".m4a": "audio/x-m4a",
        ".ogg": "audio/ogg",
        ".flac": "audio/flac",
    }
    content_type = mime_map.get(audio_path.suffix.lower(), "audio/mpeg")

    headers = {
        "Authorization": f"Bearer {settings.hf_token.strip()}",
        "Content-Type": content_type,
    }
    
    with open(audio_path, "rb") as f:
        audio_bytes = f.read()

    last_error = None
    async with httpx.AsyncClient(timeout=120.0) as client:
        for endpoint in endpoints:
            for attempt in range(3):
                try:
                    response = await client.post(endpoint, headers=headers, content=audio_bytes)
                    
                    if response.status_code == 200:
                        data = response.json()
                        text = ""
                        if isinstance(data, dict):
                            text = data.get("text", "")
                        elif isinstance(data, list):
                            text = " ".join(item.get("text", "") for item in data if isinstance(item, dict))
                            
                        text = text.strip()
                        metadata = {
                            "model": model_name,
                            "endpoint_used": endpoint,
                            "audio_size_bytes": len(audio_bytes),
                        }
                        return text, metadata
                    
                    if response.status_code == 503:
                        try:
                            data = response.json()
                            est_time = float(data.get("estimated_time", 15.0))
                        except Exception:
                            est_time = 15.0
                        sleep_seconds = min(est_time, 20.0)
                        await asyncio.sleep(sleep_seconds)
                        continue
                    
                    if response.status_code == 401:
                        raise ValueError("Invalid or unauthorized Hugging Face token (HF_TOKEN). Please check your token permissions.")
                    
                    if response.status_code == 404:
                        last_error = f"Model {model_name} not found at {endpoint}"
                        break

                    error_detail = response.text[:300]
                    last_error = f"Hugging Face API returned HTTP {response.status_code}: {error_detail}"
                    break
                except httpx.RequestError as exc:
                    last_error = f"Hugging Face network error: {str(exc)}"
                    await asyncio.sleep(2.0)
                    
    raise RuntimeError(last_error or "Transcription failed across Hugging Face endpoints.")


# =========================================================================
# MODULE 3: SUMMARIZATION, KEYWORDS & FLASHCARD SERVICES
# =========================================================================

def chunk_text(text: str, max_chars: int = 2500) -> List[str]:
    """Split text into coherent chunks of at most max_chars, breaking on sentences or paragraphs.
    Never truncates text.
    """
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        paragraphs = [text.strip()]

    chunks: List[str] = []
    current_chunk = ""

    for para in paragraphs:
        if len(para) > max_chars:
            # Paragraph itself is too big: split on sentence boundaries
            sentences = [s.strip() for s in re.split(r'(?<=[.!?])\s+', para) if s.strip()]
            for sentence in sentences:
                if len(sentence) > max_chars:
                    # Rare extreme sentence: split on words
                    words = sentence.split()
                    for word in words:
                        if len(current_chunk) + len(word) + 1 <= max_chars:
                            current_chunk += (" " if current_chunk else "") + word
                        else:
                            if current_chunk:
                                chunks.append(current_chunk)
                            current_chunk = word
                else:
                    if len(current_chunk) + len(sentence) + 1 <= max_chars:
                        current_chunk += ("\n" if current_chunk else "") + sentence
                    else:
                        if current_chunk:
                            chunks.append(current_chunk)
                        current_chunk = sentence
        else:
            if len(current_chunk) + len(para) + 2 <= max_chars:
                current_chunk += ("\n\n" if current_chunk else "") + para
            else:
                if current_chunk:
                    chunks.append(current_chunk)
                current_chunk = para

    if current_chunk:
        chunks.append(current_chunk)

    return chunks if chunks else [text]


async def query_bart_summary(
    client: httpx.AsyncClient,
    text: str,
    max_len: int = 150,
    min_len: int = 35,
) -> str:
    """Call Hugging Face inference for facebook/bart-large-cnn with retries."""
    if not settings.hf_token or not settings.hf_token.strip():
        raise ValueError("HF_TOKEN is missing. Please set your Hugging Face API token in .env.")

    model_name = settings.hf_summary_model or "facebook/bart-large-cnn"
    endpoints = [
        f"{HF_ROUTER_URL}/{model_name}",
        f"{HF_LEGACY_URL}/{model_name}",
    ]
    headers = {
        "Authorization": f"Bearer {settings.hf_token.strip()}",
        "Content-Type": "application/json",
    }

    # Ensure min/max length parameters are valid relative to input words
    input_word_count = len(text.split())
    eff_min_len = min(min_len, max(5, int(input_word_count * 0.4)))
    eff_max_len = max(eff_min_len + 15, min(max_len, max(30, int(input_word_count * 0.9))))

    payload = {
        "inputs": text,
        "parameters": {
            "max_length": eff_max_len,
            "min_length": eff_min_len,
            "do_sample": False,
        },
    }

    last_error = None
    for endpoint in endpoints:
        for attempt in range(3):
            try:
                res = await client.post(endpoint, headers=headers, json=payload, timeout=60.0)
                if res.status_code == 200:
                    data = res.json()
                    if isinstance(data, list) and len(data) > 0:
                        return data[0].get("summary_text", "").strip()
                    elif isinstance(data, dict):
                        return data.get("summary_text", "").strip()
                    return ""
                if res.status_code == 503:
                    # Model booting up
                    await asyncio.sleep(5.0)
                    continue
                if res.status_code == 401:
                    raise ValueError("Unauthorized HF_TOKEN. Please verify your Hugging Face token permissions.")
                last_error = f"BART API error ({res.status_code}): {res.text[:200]}"
                break
            except httpx.RequestError as e:
                last_error = f"Network error contacting BART: {str(e)}"
                await asyncio.sleep(2.0)

    raise RuntimeError(last_error or "BART summarization failed across Hugging Face endpoints.")


async def summarize_text_hf(text: str) -> Tuple[str, List[str], int]:
    """Chunk and hierarchically summarize text with facebook/bart-large-cnn, never silently truncating.
    Returns: (summary_text, key_points, chunks_processed)
    """
    clean_text = text.strip()
    if not clean_text:
        return "No text available to summarize.", [], 0

    # Short texts can be summarized directly
    if len(clean_text) <= 2500:
        async with httpx.AsyncClient(timeout=60.0) as client:
            summary = await query_bart_summary(client, clean_text, max_len=200, min_len=40)
            key_points = [
                s.strip() for s in re.split(r'(?<=[.!?])\s+', summary)
                if len(s.strip()) > 15
            ]
            return summary, key_points, 1

    # Text exceeds model input limit: Chunk -> Summarize each chunk -> Combine & synthesize
    chunks = chunk_text(clean_text, max_chars=2500)
    intermediate_summaries = []

    async with httpx.AsyncClient(timeout=90.0) as client:
        for chunk in chunks:
            chunk_summary = await query_bart_summary(client, chunk, max_len=130, min_len=30)
            if chunk_summary:
                intermediate_summaries.append(chunk_summary)

        combined_intermediate = " ".join(intermediate_summaries)

        # If combined intermediate summaries still exceed 2500 chars, chunk hierarchically
        if len(combined_intermediate) > 2500:
            second_chunks = chunk_text(combined_intermediate, max_chars=2500)
            second_summaries = []
            for sc in second_chunks:
                ss = await query_bart_summary(client, sc, max_len=120, min_len=30)
                if ss:
                    second_summaries.append(ss)
            combined_intermediate = " ".join(second_summaries)

        final_summary = await query_bart_summary(client, combined_intermediate, max_len=250, min_len=60)

    key_points = [
        s.strip() for s in re.split(r'(?<=[.!?])\s+', final_summary)
        if len(s.strip()) > 15
    ]
    return final_summary, key_points, len(chunks)


def extract_keywords_tfidf(text: str, top_n: int = 15) -> List[Dict[str, Any]]:
    """Extract key terms/concepts using TF-IDF with custom academic stopword filtering.
    Returns: List of dicts with keys: term, score, frequency.
    """
    clean_text = text.strip()
    if not clean_text:
        return []

    # Split into sentences to treat each sentence as a sub-document
    sentences = [s.strip() for s in re.split(r'[.\n!?]+', clean_text) if len(s.strip()) > 10]
    if len(sentences) <= 1:
        sentences = [clean_text]

    try:
        vec = TfidfVectorizer(
            stop_words=CUSTOM_STOP_WORDS,
            ngram_range=(1, 2),
            token_pattern=r'(?u)\b[a-zA-Z]{3,}\b',
            max_features=50,
        )
        X = vec.fit_transform(sentences)
        scores = X.mean(axis=0).A1
        feature_names = vec.get_feature_names_out()

        lowered = clean_text.lower()
        keywords = []
        for idx in scores.argsort()[::-1]:
            term = feature_names[idx]
            score = round(float(scores[idx]), 3)
            if score <= 0.02:
                continue
            freq = len(re.findall(r'\b' + re.escape(term) + r'\b', lowered))
            keywords.append({
                "term": term.title(),
                "score": score,
                "frequency": max(freq, 1),
            })
            if len(keywords) >= top_n:
                break
        return keywords
    except Exception as e:
        print(f"TF-IDF keyword extraction error: {e}")
        return []


def validate_qa_pair(card: Dict[str, Any], text_context: str) -> bool:
    """Robustly validate that question and answer are non-empty, related, and NEVER MCQs."""
    if not isinstance(card, dict):
        return False

    q = str(card.get("question", "")).strip()
    a = str(card.get("answer", "")).strip()

    # Minimum acceptable length
    if len(q) < 8 or len(a) < 3:
        return False

    # Strictly reject multiple-choice questions (never fall back to MCQs)
    mcq_pattern = r'(\b[A-D]\s*[\)\.:]|\bOption\s+[A-D]\b|\bWhich of the following\b|\bSelect the correct\b|\(A\)\s|\(B\)\s)'
    if re.search(mcq_pattern, q, re.IGNORECASE) or re.search(mcq_pattern, a, re.IGNORECASE):
        return False

    # Reject trivial placeholders
    placeholders = {"n/a", "none", "unknown", "tbd", "undefined", "true", "false", "yes", "no"}
    if a.lower() in placeholders:
        return False

    # Relatedness check: Q & A share vocabulary or connect to the lecture context
    q_words = set(re.findall(r'\b[a-zA-Z]{3,}\b', q.lower())) - set(ENGLISH_STOP_WORDS)
    a_words = set(re.findall(r'\b[a-zA-Z]{3,}\b', a.lower())) - set(ENGLISH_STOP_WORDS)
    ctx_words = set(re.findall(r'\b[a-zA-Z]{3,}\b', text_context.lower())) - set(ENGLISH_STOP_WORDS)

    shares_qa = bool(q_words & a_words)
    relates_to_ctx = bool((q_words | a_words) & ctx_words)

    return shares_qa or relates_to_ctx


async def generate_flashcards_hf(
    summary: str,
    keywords: List[Dict[str, Any]],
    count: int = 10,
) -> List[Dict[str, Any]]:
    """Generate high-yield Q&A flashcards using meta-llama/Llama-3.1-8B-Instruct via HF Router.
    Validates every card and skips/rejects malformed cards or MCQs.
    """
    if not settings.hf_token or not settings.hf_token.strip():
        raise ValueError("HF_TOKEN is missing. Please set your Hugging Face API token in .env.")

    kw_list = [k["term"] for k in keywords[:10]]
    kw_str = ", ".join(kw_list) if kw_list else "General lecture content"

    prompt = f"""You are an expert academic tutor creating flashcards for students.
Create exactly {count} high-yield study flashcards based on the lecture summary and key concepts below.

RULES:
1. Every card must have a clear "question" testing a concept, mechanism, definition, or distinction.
2. Every card must have a concise, accurate "answer" (1-3 sentences).
3. NEVER generate multiple choice questions (no options like A, B, C, D, no "Which of the following").
4. Both question and answer must be substantive and directly related to the material.
5. Provide a relevant academic "category" for each card (e.g. "Definition", "Architecture", "Key Concept").
6. Output ONLY a valid JSON array of objects with no surrounding conversation or markdown outside the array.

Format:
[
  {{
    "question": "What is ...?",
    "answer": "...",
    "category": "..."
  }}
]

Lecture Summary:
{summary}

Key Concepts:
{kw_str}
"""

    headers = {
        "Authorization": f"Bearer {settings.hf_token.strip()}",
        "Content-Type": "application/json",
    }
    payload = {
        "model": settings.hf_flashcard_model or "meta-llama/Llama-3.1-8B-Instruct",
        "messages": [
            {"role": "system", "content": "You are an expert educational study-aid assistant. You only output valid JSON arrays."},
            {"role": "user", "content": prompt},
        ],
        "temperature": 0.3,
        "max_tokens": 2000,
    }

    async with httpx.AsyncClient(timeout=90.0) as client:
        response = await client.post(HF_CHAT_COMPLETIONS_URL, headers=headers, json=payload)
        if response.status_code != 200:
            raise RuntimeError(f"Flashcard generation failed ({response.status_code}): {response.text[:200]}")

        data = response.json()
        content = data["choices"][0]["message"]["content"].strip()

        # Clean markdown fences if model returned them
        clean_content = re.sub(r"^```(?:json)?\s*", "", content, flags=re.MULTILINE)
        clean_content = re.sub(r"\s*```$", "", clean_content, flags=re.MULTILINE).strip()

        raw_cards = []
        try:
            raw_cards = json.loads(clean_content)
        except json.JSONDecodeError:
            match = re.search(r'\[\s*\{.*\}\s*\]', clean_content, re.DOTALL)
            if match:
                raw_cards = json.loads(match.group(0))
            else:
                raise RuntimeError("Failed to parse valid JSON flashcard output from LLM.")

        valid_cards = []
        for card in raw_cards:
            if validate_qa_pair(card, summary):
                valid_cards.append({
                    "id": str(uuid.uuid4())[:8],
                    "question": card["question"].strip(),
                    "answer": card["answer"].strip(),
                    "category": card.get("category", "Key Concept").strip(),
                })

        return valid_cards


async def generate_flashcards_for_lecture(
    lecture_id: str,
    user_id: str,
    count: Optional[int] = None,
) -> Dict[str, Any]:
    """Regenerate flashcards on demand for a lecture scoped to its owner."""
    if settings.db is None:
        raise RuntimeError("Database unavailable")

    target_count = count or settings.default_flashcard_count or 10

    # Retrieve summary or transcript to base flashcards on
    summary_doc = await settings.db.summaries.find_one({"lecture_id": lecture_id, "user_id": user_id})
    keywords_doc = await settings.db.keywords.find_one({"lecture_id": lecture_id, "user_id": user_id})

    summary_text = ""
    if summary_doc:
        summary_text = summary_doc.get("summary_text", "")
    else:
        transcript_doc = await settings.db.transcripts.find_one({"lecture_id": lecture_id, "user_id": user_id})
        if transcript_doc:
            summary_text = transcript_doc.get("text", "")[:3000]

    if not summary_text.strip():
        raise ValueError("Cannot generate flashcards: No summary or transcript found for this lecture.")

    keywords_list = keywords_doc.get("keywords", []) if keywords_doc else []

    cards = await generate_flashcards_hf(summary_text, keywords_list, count=target_count)

    flashcards_doc = {
        "lecture_id": lecture_id,
        "user_id": user_id,
        "cards": cards,
        "total_cards": len(cards),
        "model": settings.hf_flashcard_model,
        "created_at": datetime.utcnow(),
        "updated_at": datetime.utcnow(),
    }

    res = await settings.db.flashcards.update_one(
        {"lecture_id": lecture_id, "user_id": user_id},
        {"$set": flashcards_doc},
        upsert=True,
    )

    saved_doc = await settings.db.flashcards.find_one({"lecture_id": lecture_id, "user_id": user_id})
    saved_doc["id"] = str(saved_doc["_id"])
    return saved_doc


# =========================================================================
# LECTURE STATUS & PIPELINE ORCHESTRATION
# =========================================================================

async def update_lecture_status(
    lecture_id: str,
    status: str,
    status_message: Optional[str] = None,
    error_message: Optional[str] = None,
) -> None:
    """Helper to update a lecture's processing status in MongoDB."""
    if settings.db is None:
        return

    update_fields: Dict[str, Any] = {
        "processing_status": status,
        "updated_at": datetime.utcnow(),
    }
    if status_message is not None:
        update_fields["status_message"] = status_message
    if error_message is not None:
        update_fields["error_message"] = error_message
    elif status != "failed":
        update_fields["error_message"] = None

    try:
        oid = ObjectId(lecture_id)
        await settings.db.lectures.update_one({"_id": oid}, {"$set": update_fields})
    except Exception as e:
        print(f"Error updating lecture status for {lecture_id}: {e}")


async def process_lecture_background(lecture_id: str, user_id: str) -> None:
    """Orchestrates lecture processing through the full end-to-end pipeline:
    extract/transcribe -> summarize -> extract keywords -> generate flashcards -> completed.
    """
    if settings.db is None:
        print(f"Cannot process lecture {lecture_id}: Database handle is None")
        return

    try:
        oid = ObjectId(lecture_id)
    except Exception:
        return

    lecture_doc = await settings.db.lectures.find_one({"_id": oid, "user_id": user_id})
    if not lecture_doc:
        print(f"Lecture {lecture_id} not found for user {user_id}")
        return

    storage_path = Path(lecture_doc.get("storage_path", ""))
    if not storage_path.is_file():
        await update_lecture_status(
            lecture_id,
            status="failed",
            status_message="File missing on disk",
            error_message="Uploaded lecture source file could not be found on server storage.",
        )
        return

    ext = storage_path.suffix.lower()
    temp_audio_to_delete: Optional[Path] = None

    try:
        full_text = ""
        metadata: Dict[str, Any] = {}
        source_type = "document"

        # ---------------------------------------------------------
        # STEP 1: EXTRACT / TRANSCRIBE
        # ---------------------------------------------------------
        if ext in DOCUMENT_EXTENSIONS:
            await update_lecture_status(
                lecture_id,
                status="extracting",
                status_message=f"Extracting content from {ext.upper().replace('.', '')} document...",
            )
            await asyncio.sleep(0.3)

            if ext == ".pdf":
                full_text, metadata = extract_text_from_pdf(storage_path)
            elif ext in {".docx", ".doc"}:
                full_text, metadata = extract_text_from_docx(storage_path)
            elif ext in {".txt", ".md"}:
                full_text, metadata = extract_text_from_txt(storage_path)

            source_type = "document"

            if not full_text.strip():
                full_text = "[No selectable text could be extracted from this document.]"
                metadata["warning"] = "empty_or_scanned_document"

        elif ext in AUDIO_EXTENSIONS or ext in VIDEO_EXTENSIONS:
            audio_path = storage_path

            if ext in VIDEO_EXTENSIONS:
                await update_lecture_status(
                    lecture_id,
                    status="extracting",
                    status_message=f"Extracting audio track from {ext.upper().replace('.', '')} video...",
                )
                await asyncio.sleep(0.3)
                audio_path = extract_audio_from_video(storage_path)
                temp_audio_to_delete = audio_path

            await update_lecture_status(
                lecture_id,
                status="transcribing",
                status_message=f"Transcribing audio using Whisper AI ({settings.hf_whisper_model})...",
            )
            await asyncio.sleep(0.3)

            full_text, metadata = await transcribe_audio_hf(audio_path)
            source_type = "audio"

            if not full_text.strip():
                full_text = "[No speech was recognized in the audio track.]"
                metadata["warning"] = "no_speech_detected"

        else:
            raise ValueError(f"Unsupported file format for processing: {ext}")

        # Store transcript in "transcripts" collection
        word_count = len(full_text.split())
        char_count = len(full_text)

        transcript_doc = {
            "lecture_id": lecture_id,
            "user_id": user_id,
            "text": full_text,
            "source_type": source_type,
            "word_count": word_count,
            "character_count": char_count,
            "metadata": metadata,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow(),
        }

        await settings.db.transcripts.update_one(
            {"lecture_id": lecture_id, "user_id": user_id},
            {"$set": transcript_doc},
            upsert=True,
        )

        # Check if text is meaningful enough for AI summarization & flashcards
        if word_count < 10 or metadata.get("warning"):
            await update_lecture_status(
                lecture_id,
                status="completed",
                status_message=f"Text extracted ({word_count} words). Note: Limited text available for deep AI analysis.",
            )
            return

        # ---------------------------------------------------------
        # STEP 2: SUMMARIZE (BART)
        # ---------------------------------------------------------
        await update_lecture_status(
            lecture_id,
            status="summarizing",
            status_message=f"Generating hierarchical summary with BART AI ({settings.hf_summary_model})...",
        )
        await asyncio.sleep(0.3)

        summary_text, key_points, chunks_processed = await summarize_text_hf(full_text)

        summary_doc = {
            "lecture_id": lecture_id,
            "user_id": user_id,
            "summary_text": summary_text,
            "key_points": key_points,
            "word_count": len(summary_text.split()),
            "character_count": len(summary_text),
            "chunks_processed": chunks_processed,
            "model": settings.hf_summary_model,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow(),
        }

        await settings.db.summaries.update_one(
            {"lecture_id": lecture_id, "user_id": user_id},
            {"$set": summary_doc},
            upsert=True,
        )

        # ---------------------------------------------------------
        # STEP 3: EXTRACT KEYWORDS / CONCEPTS (TF-IDF)
        # ---------------------------------------------------------
        await update_lecture_status(
            lecture_id,
            status="extracting_keywords",
            status_message="Extracting key concepts and terms via TF-IDF analysis...",
        )
        await asyncio.sleep(0.3)

        keywords = extract_keywords_tfidf(full_text, top_n=15)

        keywords_doc = {
            "lecture_id": lecture_id,
            "user_id": user_id,
            "keywords": keywords,
            "method": "TF-IDF (scikit-learn)",
            "total_keywords": len(keywords),
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow(),
        }

        await settings.db.keywords.update_one(
            {"lecture_id": lecture_id, "user_id": user_id},
            {"$set": keywords_doc},
            upsert=True,
        )

        # ---------------------------------------------------------
        # STEP 4: GENERATE FLASHCARDS (LLAMA 3.1)
        # ---------------------------------------------------------
        target_card_count = settings.default_flashcard_count or 10
        await update_lecture_status(
            lecture_id,
            status="generating_flashcards",
            status_message=f"Generating {target_card_count} study flashcards with Llama 3.1 AI...",
        )
        await asyncio.sleep(0.3)

        cards = await generate_flashcards_hf(summary_text, keywords, count=target_card_count)

        flashcards_doc = {
            "lecture_id": lecture_id,
            "user_id": user_id,
            "cards": cards,
            "total_cards": len(cards),
            "model": settings.hf_flashcard_model,
            "created_at": datetime.utcnow(),
            "updated_at": datetime.utcnow(),
        }

        await settings.db.flashcards.update_one(
            {"lecture_id": lecture_id, "user_id": user_id},
            {"$set": flashcards_doc},
            upsert=True,
        )

        # ---------------------------------------------------------
        # STEP 5: MARK COMPLETED
        # ---------------------------------------------------------
        await update_lecture_status(
            lecture_id,
            status="completed",
            status_message=f"Processing complete: Summary, {len(keywords)} concepts, and {len(cards)} flashcards ready.",
        )
        print(f" Lecture {lecture_id} fully processed through all Module 3 stages.")

    except Exception as exc:
        error_msg = str(exc)
        print(f"❌ Processing failed for lecture {lecture_id}: {error_msg}")
        await update_lecture_status(
            lecture_id,
            status="failed",
            status_message="Processing failed",
            error_message=error_msg,
        )

    finally:
        # Clean up temporary extracted audio file
        if temp_audio_to_delete and temp_audio_to_delete.exists():
            try:
                temp_audio_to_delete.unlink()
                print(f" Cleaned up temporary audio: {temp_audio_to_delete.name}")
            except Exception as e:
                print(f"Notice: Failed to delete temp audio file: {e}")
